"""Tabular findings exports — CSV (stdlib) + XLSX (openpyxl).

One row per finding with every column a Burp Enterprise xlsx carries — Name, Severity, Confidence,
Location, Request, Response, Issue Detail/Background, Remediation, CWE — PLUS columns Burp's export
does NOT have: Verified (deterministic replay), Payload, Repro (curl), Detection method, Param,
OWASP. A Summary sheet gives the severity/category rollup; an optional Burp-comparison sheet renders
the coverage diff (BOTH / d4st-only / Burp-only) so the two tools sit side-by-side in one workbook.
"""

from __future__ import annotations

import csv as _csv

from .report import SEV_RANK, _meta_for

# (header, column-width) — order is the row order too.
_COLUMNS = [
    ("#", 5), ("Name", 34), ("Category", 20), ("Severity", 12), ("Confidence", 12),
    ("Verified", 11), ("Tool", 11), ("URL / Location", 58), ("Method", 8), ("Param", 16),
    ("Payload", 26), ("Description", 60), ("Remediation", 60), ("CWE", 16), ("OWASP", 26),
    ("Detection", 22), ("Request", 55), ("Response", 55), ("Repro (curl)", 48),
]
_HEADERS = [c for c, _ in _COLUMNS]

_SEV_FILL = {  # deliverable severity palette (ARGB)
    "critical": "6B21A8", "high": "B91C1C", "medium": "C2410C",
    "low": "1D4ED8", "info": "475569",
}


def _verified_str(v) -> str:
    return {True: "CONFIRMED", False: "refuted", None: "unverified"}.get(v, str(v))


def _fmt_side(d: dict) -> str:
    if not isinstance(d, dict):
        return ""
    parts: list[str] = []
    line = f"{d.get('method', '')} {d.get('url', '')}".strip()
    if line:
        parts.append(line)
    if d.get("status") is not None:
        parts.append(f"HTTP {d.get('status')}")
    for k, val in list((d.get("headers") or {}).items())[:12]:
        parts.append(f"{k}: {val}")
    if d.get("body"):
        parts.append("")
        parts.append(str(d.get("body"))[:6000])
    return "\n".join(parts)


def _exchange_text(evlog) -> tuple[str, str]:
    """Pull the first labeled request/response out of evidence_log into (request, response) text."""
    if not isinstance(evlog, list):
        return "", ""
    for ex in evlog:
        if isinstance(ex, dict) and (ex.get("request") or ex.get("response")):
            return _fmt_side(ex.get("request")), _fmt_side(ex.get("response"))
    for ex in evlog:  # fallback: JS snippet / other evidence blobs
        if isinstance(ex, dict) and ex.get("snippet"):
            return "", str(ex["snippet"])[:6000]
        if isinstance(ex, dict) and ex.get("affected_endpoints"):
            return "", "affects: " + ", ".join(ex["affected_endpoints"][:50])
    return "", ""


def _ordered(findings: list[dict]) -> list[dict]:
    return sorted(findings, key=lambda f: (SEV_RANK.get(_meta_for(f.get("category", "other"))["severity"], 9),
                                           f.get("category", ""), f.get("url", "")))


def _row(i: int, f: dict) -> list:
    m = _meta_for(f.get("category", "other"))
    req, resp = _exchange_text(f.get("evidence_log"))
    desc = (m["desc"] + ((" — " + f["evidence"]) if f.get("evidence") else "")).strip()
    return [i, m["title"], f.get("category", ""), m["severity"], f.get("confidence") or "",
            _verified_str(f.get("verified")), f.get("tool", ""), f.get("url", ""),
            f.get("method", "GET"), f.get("param") or "", f.get("payload") or "", desc,
            m["fix"], m["cwe"], m["owasp"], f.get("detection") or "", req, resp, f.get("repro") or ""]


def to_csv(result: dict, path: str) -> int:
    findings = result.get("findings") or []
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = _csv.writer(fh)
        w.writerow(_HEADERS)
        for i, f in enumerate(_ordered(findings), 1):
            w.writerow(_row(i, f))
    return len(findings)


def to_xlsx(result: dict, path: str, meta: dict | None = None, burp_diff: dict | None = None) -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from collections import Counter

    findings = _ordered(result.get("findings") or [])
    meta = meta or {}
    wb = openpyxl.Workbook()

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="0F172A")
    wrap_top = Alignment(vertical="top", wrap_text=True)

    # ---- Summary sheet -------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    sev_counts = Counter(_meta_for(f.get("category", "other"))["severity"] for f in findings)
    cat_counts = Counter(f.get("category", "?") for f in findings)
    ws["A1"] = "d4st Engagement — Findings Summary"
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    ws.cell(r, 1, "Target").font = Font(bold=True)
    ws.cell(r, 2, meta.get("target") or result.get("target") or "")
    r += 1
    ws.cell(r, 1, "Client").font = Font(bold=True)
    ws.cell(r, 2, meta.get("client") or "")
    r += 1
    ws.cell(r, 1, "Total findings").font = Font(bold=True)
    ws.cell(r, 2, len(findings))
    r += 2
    ws.cell(r, 1, "By severity").font = Font(bold=True)
    r += 1
    for sev in ("critical", "high", "medium", "low", "info"):
        if sev_counts.get(sev):
            c = ws.cell(r, 1, sev.upper())
            c.fill = PatternFill("solid", fgColor=_SEV_FILL[sev])
            c.font = Font(bold=True, color="FFFFFF")
            ws.cell(r, 2, sev_counts[sev])
            r += 1
    r += 1
    ws.cell(r, 1, "By category").font = Font(bold=True)
    r += 1
    for cat, n in cat_counts.most_common():
        ws.cell(r, 1, cat)
        ws.cell(r, 2, n)
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16

    # ---- Findings sheet ------------------------------------------------------------------------
    fs = wb.create_sheet("Findings")
    for ci, (h, w) in enumerate(_COLUMNS, 1):
        cell = fs.cell(1, ci, h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        fs.column_dimensions[get_column_letter(ci)].width = w
    for i, f in enumerate(findings, 1):
        row = _row(i, f)
        for ci, val in enumerate(row, 1):
            cell = fs.cell(i + 1, ci, val)
            cell.alignment = wrap_top
        sev = row[3]
        sc = fs.cell(i + 1, 4)
        sc.fill = PatternFill("solid", fgColor=_SEV_FILL.get(sev, "475569"))
        sc.font = Font(bold=True, color="FFFFFF")
    fs.freeze_panes = "A2"
    fs.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{len(findings) + 1}"

    # ---- Burp comparison sheet (optional) ------------------------------------------------------
    if burp_diff:
        cs = wb.create_sheet("vs Burp")
        cs["A1"] = "d4st vs Burp — coverage (same target)"
        cs["A1"].font = Font(bold=True, size=13)
        hdrs = ["Bucket", "Category", "Endpoint", "Param", "Severity"]
        for ci, h in enumerate(hdrs, 1):
            c = cs.cell(3, ci, h)
            c.font = hdr_font
            c.fill = hdr_fill
        rr = 4
        for bucket in ("both", "d4st_only", "burp_only"):
            for item in burp_diff.get(bucket, []):
                cs.cell(rr, 1, {"both": "BOTH", "d4st_only": "D4ST-ONLY",
                                "burp_only": "BURP-ONLY"}[bucket])
                cs.cell(rr, 2, item.get("category", ""))
                cs.cell(rr, 3, item.get("endpoint", ""))
                cs.cell(rr, 4, item.get("param", ""))
                cs.cell(rr, 5, item.get("severity", ""))
                rr += 1
        for ci, w in enumerate((14, 22, 52, 18, 12), 1):
            cs.column_dimensions[get_column_letter(ci)].width = w
        cs.freeze_panes = "A4"

    wb.save(path)
    return len(findings)
